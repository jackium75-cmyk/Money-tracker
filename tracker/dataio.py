"""Import from / export to the Money Tracker workbook layout, plus CSV."""
import csv
import io
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction as db_transaction

from . import services
from .models import Budget, Category, Income, IncomeSource, PaymentMethod, Transaction

EXCEL_EPOCH = date(1899, 12, 30)


class ImportProblem(Exception):
    pass


# --------------------------------------------------------------------------- cell parsing

def to_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return EXCEL_EPOCH + timedelta(days=int(v))
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def to_time(v):
    if v in (None, ""):
        return None
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, (int, float)):
        secs = round((float(v) % 1) * 86400)
        return time(secs // 3600 % 24, secs // 60 % 60)
    s = str(v).strip().upper()
    for f in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p", "%I:%M%p"):
        try:
            return datetime.strptime(s, f).time()
        except ValueError:
            continue
    return None


def to_amount(v):
    if v in (None, ""):
        return None
    try:
        d = Decimal(str(v).replace(",", "").replace("KES", "").strip())
    except InvalidOperation:
        return None
    return d.quantize(Decimal("0.01")) if d > 0 else None


def text(v):
    return "" if v is None else str(v).strip()


class Lookup:
    """Finds-or-creates list items by (case-insensitive) name."""

    def __init__(self, user, model, **create_defaults):
        self.user, self.model, self.defaults = user, model, create_defaults
        self.cache = {o.name.lower(): o for o in model.objects.filter(user=user)}
        self.created = 0

    def get(self, name):
        name = text(name)
        if not name or name.startswith("="):
            return None
        obj = self.cache.get(name.lower())
        if obj is None:
            obj = self.model.objects.create(
                user=self.user, name=name[:80], order=len(self.cache) + 1, **self.defaults
            )
            self.cache[name.lower()] = obj
            self.created += 1
        return obj


# --------------------------------------------------------------------------- workbook import

def _resolve(ws_by_name, v):
    """Budget!A5 is '=Settings!$B$6' - follow simple sheet references when no cached value exists."""
    if isinstance(v, str) and v.startswith("="):
        m = re.fullmatch(r"=\s*'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", v.strip())
        if m and m.group(1) in ws_by_name:
            return ws_by_name[m.group(1)][f"{m.group(2)}{m.group(3)}"].value
        return None
    return v


@db_transaction.atomic
def import_workbook(user, fileobj):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        raise ImportProblem("openpyxl is not installed on the server")
    try:
        wb_raw = openpyxl.load_workbook(fileobj, data_only=False)
        fileobj.seek(0)
        wb_val = openpyxl.load_workbook(fileobj, data_only=True)
    except Exception:
        raise ImportProblem("That doesn't look like an Excel (.xlsx) file")

    raw = {ws.title: ws for ws in wb_raw}
    val = {ws.title: ws for ws in wb_val}
    if not {"Transactions", "Income", "Budget", "Settings"} & set(val):
        raise ImportProblem("No Transactions, Income, Budget or Settings sheet found in this workbook")

    def cell(sheet, ref):
        v = val[sheet][ref].value
        return v if v is not None else _resolve(raw, raw[sheet][ref].value)

    services.seed_defaults(user)
    summary = {"transactions": 0, "income": 0, "budgets": 0, "skipped": 0, "new_list_items": 0}
    cats = Lookup(user, Category)
    methods = Lookup(user, PaymentMethod)
    sources = Lookup(user, IncomeSource)

    if "Settings" in val:
        for r in range(6, 30):
            cats.get(cell("Settings", f"B{r}"))
            methods.get(cell("Settings", f"D{r}"))
            sources.get(cell("Settings", f"E{r}"))
        warn = val["Settings"]["K16"].value
        if isinstance(warn, (int, float)) and 0.1 <= warn <= 1:
            p = services.get_profile(user)
            p.warn_pct = Decimal(str(warn)).quantize(Decimal("0.01"))
            p.save()

    if "Transactions" in val:
        ws = val["Transactions"]
        batch = []
        for row in ws.iter_rows(min_row=5, max_col=7, values_only=True):
            d, t, item, cat, amt, paid, notes = (list(row) + [None] * 7)[:7]
            d, amt = to_date(d), to_amount(amt)
            if not d or amt is None:
                if any(x not in (None, "") for x in (item, amt)):
                    summary["skipped"] += 1
                continue
            batch.append(Transaction(
                user=user, date=d, time=to_time(t), item=text(item)[:200] or "(no description)",
                category=cats.get(cat), amount=amt, payment_method=methods.get(paid),
                notes=text(notes)[:300],
            ))
        Transaction.objects.bulk_create(batch)
        summary["transactions"] = len(batch)

    if "Income" in val:
        ws = val["Income"]
        batch = []
        for row in ws.iter_rows(min_row=5, max_col=6, values_only=True):
            d, src, desc, amt, via, notes = (list(row) + [None] * 6)[:6]
            d, amt = to_date(d), to_amount(amt)
            if not d or amt is None:
                if amt not in (None, ""):
                    summary["skipped"] += 1
                continue
            batch.append(Income(
                user=user, date=d, source=sources.get(src), description=text(desc)[:200],
                amount=amt, received_via=methods.get(via), notes=text(notes)[:300],
            ))
        Income.objects.bulk_create(batch)
        summary["income"] = len(batch)

    if "Budget" in val:
        for r in range(5, 45):
            amt = to_amount(val["Budget"][f"C{r}"].value)
            cat = cats.get(cell("Budget", f"A{r}"))
            period = text(val["Budget"][f"B{r}"].value).capitalize() or Budget.MONTHLY
            if amt is None or cat is None or period not in dict(Budget.PERIODS):
                continue
            Budget.objects.create(
                user=user, category=cat, period=period, amount=amt,
                starts=to_date(val["Budget"][f"D{r}"].value),
                ends=to_date(val["Budget"][f"E{r}"].value),
            )
            summary["budgets"] += 1

    summary["new_list_items"] = cats.created + methods.created + sources.created
    services.bump_revision(user)
    return summary


# --------------------------------------------------------------------------- CSV import

def _norm(h):
    return re.sub(r"[^a-z]", "", h.lower())


@db_transaction.atomic
def import_csv(user, fileobj, kind):
    try:
        content = fileobj.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ImportProblem("Save the CSV as UTF-8 and try again")
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        raise ImportProblem("The CSV file is empty")
    cols = {_norm(h): h for h in reader.fieldnames}

    def col(row, *names):
        for n in names:
            if n in cols:
                return row.get(cols[n])
        return None

    services.seed_defaults(user)
    methods = Lookup(user, PaymentMethod)
    n = skipped = 0
    if kind == "income":
        sources = Lookup(user, IncomeSource)
        for row in reader:
            d, amt = to_date(col(row, "date")), to_amount(col(row, "amount", "amountkes"))
            if not d or amt is None:
                skipped += 1
                continue
            Income.objects.create(
                user=user, date=d, amount=amt, source=sources.get(col(row, "source", "incomesource")),
                description=text(col(row, "description"))[:200],
                received_via=methods.get(col(row, "receivedvia", "via", "method")),
                notes=text(col(row, "notes"))[:300],
            )
            n += 1
        summary = {"income": n}
    else:
        cats = Lookup(user, Category)
        for row in reader:
            d, amt = to_date(col(row, "date")), to_amount(col(row, "amount", "amountkes"))
            if not d or amt is None:
                skipped += 1
                continue
            Transaction.objects.create(
                user=user, date=d, time=to_time(col(row, "time")), amount=amt,
                item=text(col(row, "item", "itemproduct", "description", "product"))[:200] or "(no description)",
                category=cats.get(col(row, "category")),
                payment_method=methods.get(col(row, "paidwith", "paymentmethod", "method")),
                notes=text(col(row, "notes"))[:300],
            )
            n += 1
        summary = {"transactions": n}
    summary["skipped"] = skipped
    services.bump_revision(user)
    return summary


# --------------------------------------------------------------------------- export

def transactions_csv(user, qs=None):
    qs = qs if qs is not None else Transaction.objects.filter(user=user)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Date", "Time", "Item / Product", "Category", "Amount", "Paid With", "Notes"])
    for t in qs.select_related("category", "payment_method").order_by("date", "time", "id"):
        w.writerow([
            t.date.isoformat(), t.time.strftime("%H:%M") if t.time else "", t.item,
            t.category.name if t.category else "", t.amount,
            t.payment_method.name if t.payment_method else "", t.notes,
        ])
    return out.getvalue()


def income_csv(user):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Date", "Income Source", "Description", "Amount", "Received Via", "Notes"])
    for i in Income.objects.filter(user=user).select_related("source", "received_via").order_by("date", "id"):
        w.writerow([
            i.date.isoformat(), i.source.name if i.source else "", i.description, i.amount,
            i.received_via.name if i.received_via else "", i.notes,
        ])
    return out.getvalue()


def workbook_bytes(user):
    """A backup in the same sheet layout as the Money Tracker workbook (re-importable)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    cur = services.get_profile(user).currency
    wb = openpyxl.Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E3A5F")

    def sheet(title, subtitle, headers, widths):
        ws = wb.create_sheet(title)
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="666666")
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = widths[c - 1]
        ws.freeze_panes = "A5"
        return ws

    wb.remove(wb.active)
    ws = sheet("Transactions", f"Exported {services.today():%d %b %Y}",
               ["Date", "Time", "Item / Product", "Category", f"Amount ({cur})", "Paid With", "Notes"],
               [12, 10, 34, 24, 14, 16, 30])
    for r, t in enumerate(
        Transaction.objects.filter(user=user).select_related("category", "payment_method").order_by("date", "time", "id"), 5
    ):
        ws.cell(r, 1, t.date).number_format = "dd mmm yyyy"
        if t.time:
            ws.cell(r, 2, t.time).number_format = "h:mm AM/PM"
        ws.cell(r, 3, t.item)
        ws.cell(r, 4, t.category.name if t.category else "")
        ws.cell(r, 5, float(t.amount)).number_format = "#,##0.00"
        ws.cell(r, 6, t.payment_method.name if t.payment_method else "")
        ws.cell(r, 7, t.notes)

    ws = sheet("Income", f"Exported {services.today():%d %b %Y}",
               ["Date", "Income Source", "Description", f"Amount ({cur})", "Received Via", "Notes"],
               [12, 26, 34, 14, 16, 30])
    for r, i in enumerate(
        Income.objects.filter(user=user).select_related("source", "received_via").order_by("date", "id"), 5
    ):
        ws.cell(r, 1, i.date).number_format = "dd mmm yyyy"
        ws.cell(r, 2, i.source.name if i.source else "")
        ws.cell(r, 3, i.description)
        ws.cell(r, 4, float(i.amount)).number_format = "#,##0.00"
        ws.cell(r, 5, i.received_via.name if i.received_via else "")
        ws.cell(r, 6, i.notes)

    ws = sheet("Budget", "Daily, weekly or monthly budgets per category",
               ["Category", "Budget Period", f"Budget Amount ({cur})", "Starts (optional)", "Ends (optional)",
                "Per Day", "Per Week", "Per Month (avg)"], [26, 14, 18, 16, 16, 12, 12, 16])
    for r, b in enumerate(Budget.objects.filter(user=user).select_related("category"), 5):
        pd = b.per_day()
        ws.cell(r, 1, b.category.name)
        ws.cell(r, 2, b.period)
        ws.cell(r, 3, float(b.amount)).number_format = "#,##0"
        if b.starts:
            ws.cell(r, 4, b.starts).number_format = "dd mmm yyyy"
        if b.ends:
            ws.cell(r, 5, b.ends).number_format = "dd mmm yyyy"
        for c, v in ((6, pd), (7, pd * 7), (8, pd * 365 / 12)):
            ws.cell(r, c, round(v, 2)).number_format = "#,##0"

    ws = wb.create_sheet("Settings")
    ws["B1"] = "Settings & Lists"
    ws["B1"].font = Font(bold=True, size=16)
    for col, title, items in (
        ("B", "Expense Categories", Category.objects.filter(user=user)),
        ("D", "Payment Methods", PaymentMethod.objects.filter(user=user)),
        ("E", "Income Sources", IncomeSource.objects.filter(user=user)),
    ):
        ws[f"{col}5"] = title
        ws[f"{col}5"].font = Font(bold=True)
        ws.column_dimensions[col].width = 28
        for r, o in enumerate(items, 6):
            ws[f"{col}{r}"] = o.name
    ws["J16"] = "Warn me when budget used reaches"
    ws["K16"] = float(services.get_profile(user).warn_pct)
    ws["K16"].number_format = "0%"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
